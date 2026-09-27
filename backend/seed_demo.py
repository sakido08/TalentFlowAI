"""Create fictional demo accounts, candidate profiles, and jobs.

Run once after starting the database: python -m backend.seed_demo
Demo passwords are for classroom demos only and must be changed before real use.
"""

import os
from pathlib import Path

from docx import Document
from dotenv import load_dotenv

load_dotenv()

from .database import Base, SessionLocal, engine
from .helpers import save_candidate_data
from .models import Candidate, Job, Resume, User
from .schemas import CandidateUpdate, EducationInput, ExperienceInput
from .security import hash_password
from .services.extraction import extract_information
from .services.preprocessing import preprocess_text
from .services.resume_parser import extract_text


def main() -> None:
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        examples = [
            ("Alex Morgan", "alex@example.edu", "Python developer familiar with FastAPI, PostgreSQL, Git and Docker. Built a student project dashboard.", ["Python", "FastAPI", "PostgreSQL", "Git", "Docker"]),
            ("Jamie Chen", "jamie@example.edu", "Computer science student with Java, SQL and React experience. Worked on a course registration app.", ["Java", "SQL", "React", "Git"]),
            ("Sam Rivera", "sam@example.edu", "Junior data analyst. Uses Python, Pandas, SQL and scikit-learn for data analysis and machine learning.", ["Python", "Pandas", "SQL", "scikit-learn"]),
        ]
        accounts = [
            ("Demo Admin", "admin@talentflow.local", "AdminDemo123!", "admin"),
            ("Demo Recruiter", "recruiter@talentflow.local", "Recruiter123!", "recruiter"),
            ("Demo User", "user@talentflow.local", "UserDemo123!", "user"),
        ]
        for name, email, password, role in accounts:
            if not db.query(User).filter_by(email=email).first():
                db.add(User(name=name, email=email, password_hash=hash_password(password), role=role))
        db.flush()
        recruiter = db.query(User).filter_by(email="recruiter@talentflow.local").one()
        upload_dir = Path(os.getenv("UPLOAD_DIR", "uploads"))
        upload_dir.mkdir(parents=True, exist_ok=True)
        for name, email, summary, skills in examples:
            if db.query(Candidate).join(Resume).filter(Candidate.email == email).first():
                continue
            # Create a real DOCX sample and run the same extraction services as an upload.
            filename = f"demo-{name.lower().replace(' ', '-')}-resume.docx"
            file_path = upload_dir / filename
            document = Document()
            document.add_paragraph(name)
            document.add_paragraph(email)
            document.add_paragraph("+1 555 010 2040")
            document.add_paragraph("Address: 123 Example Street")
            document.add_paragraph("SUMMARY")
            document.add_paragraph(summary)
            document.add_paragraph("SKILLS")
            document.add_paragraph(", ".join(skills))
            document.add_paragraph("EDUCATION")
            document.add_paragraph("BS Computer Science | Example University | 2022 - 2026")
            document.add_paragraph("EXPERIENCE")
            document.add_paragraph("Campus Technology Lab | Student Developer | 2025 - 2026")
            document.add_paragraph("Built and documented a classroom software project.")
            document.save(file_path)
            original = extract_text(file_path.read_bytes(), filename)
            processed = preprocess_text(original)
            fields = extract_information(processed)
            resume = Resume(user_id=recruiter.id, filename=filename, stored_path=str(file_path),
                            extracted_text=original, processed_text=processed, status="completed", processing_seconds=0.04)
            db.add(resume)
            db.flush()
            candidate = Candidate(resume_id=resume.id)
            db.add(candidate)
            db.flush()
            fields["education"] = [EducationInput(**item) for item in fields["education"]]
            fields["experience"] = [ExperienceInput(**item) for item in fields["experience"]]
            save_candidate_data(db, candidate, CandidateUpdate(**fields))
        if not db.query(Job).first():
            db.add_all([
                Job(recruiter_id=recruiter.id, title="Junior Python Developer", description="Build small APIs and work with relational data.", required_skills="Python, FastAPI, SQL, Git"),
                Job(recruiter_id=recruiter.id, title="Entry-level Data Analyst", description="Prepare datasets and communicate useful insights.", required_skills="Python, SQL, Pandas, scikit-learn"),
            ])
        db.commit()
        print("Demo data is ready.")
        print("Admin:     admin@talentflow.local / AdminDemo123!")
        print("Recruiter: recruiter@talentflow.local / Recruiter123!")
        print("User:      user@talentflow.local / UserDemo123!")
    finally:
        db.close()


if __name__ == "__main__":
    main()
