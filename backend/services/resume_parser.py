"""Extract selectable text from PDF and DOCX files."""

from io import BytesIO
from pathlib import Path

import pymupdf as fitz
from docx import Document


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Return text from a PDF or DOCX file. Scanned PDFs need OCR and may be empty."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            return "\n".join(page.get_text("text") for page in document)
    if extension == ".docx":
        document = Document(BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        # Tables are common in resumes, so include their cell text too.
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(paragraphs)
    raise ValueError("Please upload a PDF or DOCX file.")
