from io import BytesIO

import pymupdf
from docx import Document

from backend.services.extraction import extract_information
from backend.services.preprocessing import preprocess_text
from backend.services.resume_parser import extract_text
from backend.services.similarity import (
    combine_scores,
    compare_text,
    education_signal,
    experience_signal,
    skill_comparison,
)


def test_preprocessing_preserves_contacts_and_dates():
    cleaned = preprocess_text("Jane Doe  \r\n jane@example.edu   \r\n 2024")
    assert "jane@example.edu" in cleaned
    assert "2024" in cleaned
    assert "\r" not in cleaned


def test_extracts_email_phone_name_and_skills():
    text = """Taylor Morgan
taylor@example.edu
+1 555 123 4567
SKILLS
Python, SQL, Git
EDUCATION
BS Computer Science - Example University 2022 - 2026
EXPERIENCE
Student Developer
Campus Lab
"""
    result = extract_information(text)
    assert result["full_name"] == "Taylor Morgan"
    assert result["email"] == "taylor@example.edu"
    assert "Python" in result["skills"]
    assert result["education"]
    assert result["experience"]


def test_missing_information_returns_empty_values():
    result = extract_information("A plain resume with no useful fields.")
    assert result["email"] == ""
    assert result["phone"] == ""
    assert result["skills"] == []


def test_reads_pdf_text():
    buffer = BytesIO()
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "PDF resume text")
    document.save(buffer)
    document.close()
    assert "PDF resume text" in extract_text(buffer.getvalue(), "resume.pdf")


def test_reads_docx_text_and_tables():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("DOCX resume text")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table skill: Python"
    document.save(buffer)
    assert "DOCX resume text" in extract_text(buffer.getvalue(), "resume.docx")
    assert "Table skill: Python" in extract_text(buffer.getvalue(), "resume.docx")


def test_rejects_unsupported_format():
    try:
        extract_text(b"content", "resume.txt")
    except ValueError as error:
        assert "PDF or DOCX" in str(error)
    else:
        raise AssertionError("Unsupported format should fail")


def test_tfidf_cosine_and_skill_comparison():
    score = compare_text("Python developer with SQL experience", "Python engineer with SQL experience")
    assert 0 < score <= 1
    assert compare_text("", "Python") == 0
    matched, missing = skill_comparison("Experienced with Python and SQL", "Python, SQL, React")
    assert matched == ["Python", "SQL"]
    assert missing == ["React"]


def test_default_weighted_score_moves_sbert_weight_to_tfidf():
    result = combine_scores(tfidf=0.8, skills=0.5, experience=1, education=1)
    assert result["overall_score"] == 74.0
    assert result["formula"].startswith("0.55")
    assert [item["weight"] for item in result["components"]] == [55, 30, 10, 5]


def test_sbert_weighted_score_uses_semantic_component():
    result = combine_scores(tfidf=0.8, skills=0.5, experience=1, education=1, semantic=0.9)
    assert result["overall_score"] == 78.0
    assert [item["weight"] for item in result["components"]] == [40, 30, 15, 10, 5]


def test_profile_signals_only_check_for_detected_information():
    assert experience_signal([{"start_date":"2021", "end_date":"2024"}]) == 1
    assert experience_signal([{"start_date":"", "end_date":""}]) == 0
    assert education_signal([{"degree":"BS Computer Science", "institution":""}]) == 1
    assert education_signal([]) == 0
